import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def make_word_doc(filename="MY_FINANCIAL_PORTFOLIO_STATEMENT.docx"):
    doc = docx.Document()
    
    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Title
    p = doc.add_paragraph()
    r = p.add_run("PERSONAL WEALTH & PORTFOLIO STATEMENT")
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    sub = doc.add_paragraph()
    r_sub = sub.add_run("Verified Financial Statement for Automatic Ingestion into Halo Wealth Engine")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(79, 70, 229)
    r_sub.bold = True

    doc.add_paragraph("Client Name: Advait Kubal | PAN: ABCDE1234F | Annual Salary Credit (CTC): ₹42,00,000")

    doc.add_heading("1. Liquid & Investment Assets", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Asset Class"
    hdr[1].text = "Description / Scheme"
    hdr[2].text = "Yield / CAGR"
    hdr[3].text = "Current Value"

    assets = [
        ("Mutual Funds", "Parag Parikh Flexi Cap & Nippon Small Cap", "14.5%", "₹95,00,000"),
        ("Equity", "Direct Equities via Zerodha Demat", "13.2%", "₹75,00,000"),
        ("Fixed Deposit", "Cumulative Bank Fixed Deposit", "7.25%", "₹30,00,000"),
        ("Gold", "Sovereign Gold Bonds (SGB)", "8.5%", "₹20,00,000"),
        ("NPS/PPF", "NPS Tier-1 & Public Provident Fund", "7.1%", "₹25,00,000"),
        ("Cash", "Liquid Savings Account Balance", "3.5%", "₹18,00,000"),
    ]

    for cat, desc, yld, val in assets:
        row = table.add_row().cells
        row[0].text = cat
        row[1].text = desc
        row[2].text = yld
        row[3].text = val

    doc.add_heading("2. Active Loan Liabilities", level=1)
    
    t_liab = doc.add_table(rows=1, cols=4)
    t_liab.alignment = WD_TABLE_ALIGNMENT.CENTER
    l_hdr = t_liab.rows[0].cells
    l_hdr[0].text = "Loan Type"
    l_hdr[1].text = "Lender / Account"
    l_hdr[2].text = "Interest Rate & EMI"
    l_hdr[3].text = "Outstanding Principal"

    liabs = [
        ("Home Loan", "HDFC Bank Housing Loan #HL-7788", "8.4% (EMI: ₹42,000)", "₹48,00,000"),
        ("Car Loan", "ICICI Express Auto Loan #AL-2233", "8.85% (EMI: ₹16,500)", "₹7,20,000"),
    ]

    for lt, lender, emi_info, bal in liabs:
        row = t_liab.add_row().cells
        row[0].text = lt
        row[1].text = lender
        row[2].text = emi_info
        row[3].text = bal

    doc.add_heading("3. Compensation & Cash Flow", level=1)
    doc.add_paragraph("• Annual Salary Credit (Gross CTC): ₹42,00,000\n• In-Hand Monthly Take-Home: ₹2,87,000\n• Monthly Mandatory Living Burn: ₹65,000")

    doc.save(filename)
    print(f"Generated Word document: {filename}")

def make_pdf_doc(filename="MY_FINANCIAL_PORTFOLIO_STATEMENT.pdf"):
    doc = SimpleDocTemplate(filename, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    c_primary = colors.HexColor("#0f172a")
    c_indigo = colors.HexColor("#4f46e5")
    c_emerald = colors.HexColor("#059669")
    c_red = colors.HexColor("#e11d48")

    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, textColor=c_primary)
    sub_style = ParagraphStyle('Sub', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, textColor=c_indigo)
    h2_style = ParagraphStyle('H2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, textColor=c_primary, spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, textColor=colors.HexColor("#334155"))
    bold_style = ParagraphStyle('Bold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, textColor=c_primary)
    tbl_hdr = ParagraphStyle('TH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, textColor=colors.white)

    story = [
        Paragraph("PERSONAL WEALTH &amp; PORTFOLIO STATEMENT", title_style),
        Paragraph("Verified Financial Statement for Automatic Ingestion into Halo Wealth Engine", sub_style),
        HRFlowable(width="100%", thickness=1.5, color=c_indigo, spaceBefore=4, spaceAfter=8),
        Paragraph("<b>Client Name:</b> Advait Kubal | <b>PAN:</b> ABCDE1234F | <b>Annual Salary Credit:</b> ₹42,00,000", body_style),
        Spacer(1, 8),
        Paragraph("1. Verified Investment &amp; Liquid Assets", h2_style)
    ]

    assets_data = [
        [Paragraph("<b>Asset Class</b>", tbl_hdr), Paragraph("<b>Description</b>", tbl_hdr), Paragraph("<b>Yield</b>", tbl_hdr), Paragraph("<b>Current Value</b>", tbl_hdr)],
        [Paragraph("Mutual Funds", bold_style), Paragraph("Parag Parikh Flexi Cap &amp; Nippon Small Cap", body_style), Paragraph("14.5%", body_style), Paragraph("₹95,00,000", bold_style)],
        [Paragraph("Equity", bold_style), Paragraph("Direct Equities via Zerodha Demat", body_style), Paragraph("13.2%", body_style), Paragraph("₹75,00,000", bold_style)],
        [Paragraph("Fixed Deposit", bold_style), Paragraph("Cumulative Bank Fixed Deposit", body_style), Paragraph("7.25%", body_style), Paragraph("₹30,00,000", bold_style)],
        [Paragraph("Gold", bold_style), Paragraph("Sovereign Gold Bonds (SGB)", body_style), Paragraph("8.5%", body_style), Paragraph("₹20,00,000", bold_style)],
        [Paragraph("NPS/PPF", bold_style), Paragraph("NPS Tier-1 &amp; Public Provident Fund", body_style), Paragraph("7.1%", body_style), Paragraph("₹25,00,000", bold_style)],
        [Paragraph("Cash", bold_style), Paragraph("Liquid Savings Account Balance", body_style), Paragraph("3.5%", body_style), Paragraph("₹18,00,000", bold_style)],
    ]
    t_a = Table(assets_data, colWidths=[90, 240, 70, 130])
    t_a.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_indigo),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_a)
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Active Loan Liabilities", h2_style))
    liab_data = [
        [Paragraph("<b>Loan Type</b>", tbl_hdr), Paragraph("<b>Lender &amp; Account</b>", tbl_hdr), Paragraph("<b>Interest &amp; EMI</b>", tbl_hdr), Paragraph("<b>Outstanding Principal</b>", tbl_hdr)],
        [Paragraph("Home Loan", bold_style), Paragraph("HDFC Housing Loan #HL-7788", body_style), Paragraph("8.40% • EMI: ₹42,000", body_style), Paragraph("₹48,00,000", bold_style)],
        [Paragraph("Car Loan", bold_style), Paragraph("ICICI Express Auto Loan #AL-2233", body_style), Paragraph("8.85% • EMI: ₹16,500", body_style), Paragraph("₹7,20,000", bold_style)],
    ]
    t_l = Table(liab_data, colWidths=[90, 200, 110, 130])
    t_l.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_red),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_l)
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Annual Compensation &amp; Monthly In-Hand Cash Flow", h2_style))
    inc_data = [
        [Paragraph("<b>Annual Salary Credit (Gross CTC):</b> ₹42,00,000", bold_style), Paragraph("<b>Monthly In-Hand Take-Home:</b> ₹2,87,000", bold_style)]
    ]
    t_i = Table(inc_data, colWidths=[265, 265])
    t_i.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f1f5f9")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_i)

    doc.build(story)
    print(f"Generated PDF document: {filename}")

if __name__ == '__main__':
    make_word_doc()
    make_pdf_doc()
